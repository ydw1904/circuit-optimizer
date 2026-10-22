"""Strict loaders for the two input files: routes and aircraft.

A file is accepted whole or rejected whole. The first malformed cell raises
DataError naming the file, row and column, and nothing from the file is used.

Accepted: .csv (UTF-8, comma or semicolon separated) and .xlsx (first sheet).
Row 1 is the header. Column order is free, every column is required, extra
columns are an error. Fully blank rows are ignored (Excel pads sheets with them).
"""

import csv
import io
import itertools
import math
import re
from pathlib import Path

MAX_BYTES = 5_000_000  # file size on disk
MAX_ROWS = 10_000      # rows read, blank ones included


class DataError(ValueError):
    """The input file is malformed; nothing was loaded from it."""


# (header, key used by the planner, kind[, lowest, highest])
# kinds: iata = 3 letters, text = short printable string,
#        int = whole number, num = any number
ROUTE_COLUMNS = (
    ("iata", "iata", "iata"),
    ("name", "name", "text"),
    ("distance_km", "dist", "int", 1, 20_100),
    ("category", "cat", "int", 1, 20),
    ("eco_demand", "eco_d", "int", 1, 1_000_000),
    ("bus_demand", "bus_d", "int", 0, 1_000_000),
    ("first_demand", "fir_d", "int", 0, 1_000_000),
    ("cargo_demand", "cargo_d", "int", 0, 1_000_000),
    ("route_price", "price", "int", 0, 10**12),  # 0 = unknown, left out of payback
)
AIRCRAFT_COLUMNS = (
    ("model", "model", "text"),
    ("category", "cat", "int", 1, 20),
    ("speed_kmh", "speed", "int", 1, 10_000),
    ("range_km", "range", "int", 1, 40_000),
    ("max_pax", "pax", "int", 0, 5_000),
    ("max_tonnage", "tonnage", "num", 0.1, 1_000),
    ("price", "price", "int", 0, 10**12),
)

_NUMBER = re.compile(r"-?\d+(\.\d+)?")  # no thousands separators, no 1e3, no nan
_IATA = re.compile(r"[A-Za-z]{3}")


def _blank(cell):
    return cell is None or str(cell).strip() == ""


def _convert(value, col, where):
    header, _, kind, *bounds = col
    if isinstance(value, str):
        value = value.strip()
    if _blank(value):
        raise DataError(f"{where}: '{header}' is empty")
    if kind == "iata":
        if not (isinstance(value, str) and _IATA.fullmatch(value)):
            raise DataError(f"{where}: '{header}' must be 3 letters, got {value!r}")
        return value.upper()
    if kind == "text":
        text = str(value)
        if len(text) > 60 or not text.isprintable():
            raise DataError(f"{where}: '{header}' must be printable text of at most 60 characters")
        return text
    # Excel hands booleans over as ints, and dates as datetimes: refuse both.
    ok_type = isinstance(value, (int, float)) and not isinstance(value, bool)
    if not ok_type and not (isinstance(value, str) and _NUMBER.fullmatch(value)):
        raise DataError(f"{where}: '{header}' must be a plain number, got {value!r}")
    num = float(value)
    if not math.isfinite(num):
        raise DataError(f"{where}: '{header}' must be a finite number, got {value!r}")
    lo, hi = bounds
    if not lo <= num <= hi:
        raise DataError(f"{where}: '{header}' must be between {lo:,} and {hi:,}, got {value!r}")
    if kind == "int":
        if num != int(num):
            raise DataError(f"{where}: '{header}' must be a whole number, got {value!r}")
        return int(num)
    return num


def _read_table(path):
    """Raw rows (lists of cells) of a .csv or .xlsx file, at most MAX_ROWS + 1 of them."""
    p = Path(path)
    if not p.is_file():
        raise DataError(f"{p}: file not found")
    if p.stat().st_size > MAX_BYTES:
        raise DataError(f"{p.name}: larger than {MAX_BYTES // 1_000_000} MB")
    kind = p.suffix.lower()
    if kind not in (".csv", ".xlsx"):
        raise DataError(f"{p.name}: unsupported file type {p.suffix!r} (use .csv or .xlsx)")
    try:
        if kind == ".csv":
            text = p.read_text(encoding="utf-8-sig")
            head = text.split("\n", 1)[0]
            delimiter = ";" if head.count(";") > head.count(",") else ","
            rows = csv.reader(io.StringIO(text), delimiter=delimiter)
            return list(itertools.islice(rows, MAX_ROWS + 1))
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            rows = wb.worksheets[0].iter_rows(values_only=True)
            return [list(r) for r in itertools.islice(rows, MAX_ROWS + 1)]
        finally:
            wb.close()
    except Exception as e:  # bad encoding, corrupt zip, NUL bytes: all just "malformed"
        raise DataError(f"{p.name}: cannot be read ({type(e).__name__}: {e})") from e


def _parse(path, columns, key, check=None):
    name = Path(path).name
    table = _read_table(path)
    if len(table) > MAX_ROWS:
        raise DataError(f"{name}: more than {MAX_ROWS:,} rows")
    if not table:
        raise DataError(f"{name}: file is empty")

    header = [("" if h is None else str(h)).strip().lower() for h in table[0]]
    while header and not header[-1]:  # Excel leaves empty trailing header cells
        header.pop()
    want = [c[0] for c in columns]
    problems = []
    if missing := [h for h in want if h not in header]:
        problems.append(f"missing column(s) {', '.join(missing)}")
    if extra := [h or "(blank)" for h in header if h not in want]:
        problems.append(f"unknown column(s) {', '.join(extra)}")
    if dupes := sorted({h for h in header if h and header.count(h) > 1}):
        problems.append(f"duplicated column(s) {', '.join(dupes)}")
    if problems:
        raise DataError(f"{name} header: {'; '.join(problems)}. Expected: {', '.join(want)}")

    at = {h: i for i, h in enumerate(header)}
    items, first_seen = [], {}
    for n, row in enumerate(table[1:], start=2):
        if all(_blank(c) for c in row):
            continue
        where = f"{name} row {n}"
        if not all(_blank(c) for c in row[len(header):]):
            raise DataError(f"{where}: values to the right of the last header column")
        item = {col[1]: _convert(row[at[col[0]]] if at[col[0]] < len(row) else None, col, where)
                for col in columns}
        if check and (problem := check(item)):
            raise DataError(f"{where}: {problem}")
        k = item[key].upper()
        if k in first_seen:
            raise DataError(f"{where}: duplicate {key} {item[key]!r} (first on row {first_seen[k]})")
        first_seen[k] = n
        items.append(item)
    if not items:
        raise DataError(f"{name}: no data rows")
    return items


def _route_problem(route):
    """Cross-field sanity check for one parsed route.

    Return a short message to reject the whole file, or None to accept the row.
    Every field already passed its own type and range check; this hook is for
    relationships between fields (for example, to catch two columns that were
    pasted in the wrong order). `route` uses the planner keys: iata, name, dist,
    cat, eco_d, bus_d, fir_d, cargo_d, price.
    """
    # TODO(human): decide which relationships between the demand columns must
    # hold for a route to be believable, and return a message when one is broken.
    return None


def load_routes(path):
    """Candidate routes for one hub. Raises DataError if the file is malformed."""
    return _parse(path, ROUTE_COLUMNS, "iata", _route_problem)


def load_aircraft(path):
    """Aircraft that may be planned. Raises DataError if the file is malformed."""
    fleet = _parse(path, AIRCRAFT_COLUMNS, "model")
    for ac in fleet:
        ac["alias"] = ac["model"]
    return fleet
